# EPUBを展開したディレクトリ（content.opfのあるディレクトリ）から、ルビ付きの全文Word（.docx）を生成する。
# 章はEPUBの読む順（spine）に従う。linear="no"のページ（表紙等）と、--excludeで指定したページ
# （既定: 目次の toc.xhtml）は含めない。表題はEPUBの書名（dc:title）から作る。
# 使い方: python build_word.py <展開したEPUBのディレクトリ> <出力.docx> [--exclude ファイル名 ...]
# 依存: python-docx beautifulsoup4 lxml pillow（システムPythonに入れず venv を使う）
import argparse
import re
import sys
from pathlib import Path
from bs4 import BeautifulSoup, NavigableString, Tag, Comment
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

_ap = argparse.ArgumentParser(description="EPUBからルビ付きの全文Wordを作る")
_ap.add_argument("epub_dir")
_ap.add_argument("out")
_ap.add_argument("--exclude", nargs="*", default=["toc.xhtml"])
_args = _ap.parse_args()
EPUB_DIR = Path(_args.epub_dir)
OUT_PATH = Path(_args.out)


def spine_files():
    """content.opfの読む順（spine）から、本文として書き出すファイルと書名を返す。"""
    opf = (EPUB_DIR / "content.opf").read_text(encoding="utf-8")
    items = dict(re.findall(r'<item [^>]*?id="([^"]+)"[^>]*?href="([^"]+)"', opf))
    nav = set(re.findall(r'<item [^>]*?href="([^"]+)"[^>]*?properties="[^"]*nav', opf))
    files = []
    for idref, rest in re.findall(r'<itemref idref="([^"]+)"([^>]*)>', opf):
        href = items.get(idref)
        if not href or 'linear="no"' in rest or href in nav or href in _args.exclude:
            continue
        files.append(href)
    m = re.search(r"<dc:title[^>]*>([^<]+)</dc:title>", opf)
    return files, (m.group(1) if m else "")

BASE_FONT = "游明朝"
BASE_SZ = 24
RUBY_SZ = 12


def is_hidden(tag):
    return isinstance(tag, Tag) and tag.get('aria-hidden') == 'true'


def set_run_font(run, font_name=BASE_FONT, size_half_points=BASE_SZ, bold=False):
    run.font.size = Pt(size_half_points / 2)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)


def add_plain_run(paragraph, text, bold=False):
    if not text:
        return
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)


def add_ruby_run(paragraph, base_text, rt_text, bold=False):
    if not base_text:
        return
    if not rt_text:
        add_plain_run(paragraph, base_text, bold=bold)
        return
    p = paragraph._p
    r = OxmlElement('w:r')
    ruby = OxmlElement('w:ruby')
    rubyPr = OxmlElement('w:rubyPr')
    rubyAlign = OxmlElement('w:rubyAlign')
    rubyAlign.set(qn('w:val'), 'distributeSpace')
    hps = OxmlElement('w:hps')
    hps.set(qn('w:val'), str(RUBY_SZ))
    hpsRaise = OxmlElement('w:hpsRaise')
    hpsRaise.set(qn('w:val'), str(int(BASE_SZ * 0.9)))
    hpsBaseText = OxmlElement('w:hpsBaseText')
    hpsBaseText.set(qn('w:val'), str(BASE_SZ))
    lid = OxmlElement('w:lid')
    lid.set(qn('w:val'), 'ja-JP')
    for el in (rubyAlign, hps, hpsRaise, hpsBaseText, lid):
        rubyPr.append(el)
    ruby.append(rubyPr)

    def build_text_run(text, size, bold_inner):
        rr = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:eastAsia'), BASE_FONT)
        rFonts.set(qn('w:ascii'), BASE_FONT)
        rFonts.set(qn('w:hAnsi'), BASE_FONT)
        rPr.append(rFonts)
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), str(size))
        rPr.append(sz)
        if bold_inner:
            rPr.append(OxmlElement('w:b'))
        rr.append(rPr)
        t = OxmlElement('w:t')
        t.set(qn('xml:space'), 'preserve')
        t.text = text
        rr.append(t)
        return rr

    rt_el = OxmlElement('w:rt')
    rt_el.append(build_text_run(rt_text, RUBY_SZ, False))
    ruby.append(rt_el)
    rubyBase_el = OxmlElement('w:rubyBase')
    rubyBase_el.append(build_text_run(base_text, BASE_SZ, bold))
    ruby.append(rubyBase_el)
    r.append(ruby)
    p.append(r)


def walk_inline(paragraph, node, bold=False):
    for child in node.children:
        if is_hidden(child):
            continue
        if isinstance(child, Comment):
            continue
        if isinstance(child, NavigableString):
            text = str(child)
            if text:
                add_plain_run(paragraph, text, bold=bold)
        elif isinstance(child, Tag):
            if child.name == 'ruby':
                rt_tag = child.find('rt')
                rt_text = rt_tag.get_text() if rt_tag else ''
                base_text = ''.join(
                    c for c in child.contents
                    if isinstance(c, NavigableString) and not isinstance(c, Comment)
                ).strip()
                if not base_text:
                    base_text = child.get_text()
                    if rt_text and base_text.endswith(rt_text):
                        base_text = base_text[: -len(rt_text)]
                add_ruby_run(paragraph, base_text, rt_text, bold=bold)
            elif child.name == 'rt':
                continue
            elif child.name in ('strong', 'b'):
                walk_inline(paragraph, child, bold=True)
            elif child.name == 'br':
                paragraph.add_run().add_break()
            else:
                walk_inline(paragraph, child, bold=bold)


def add_image(doc, img_tag):
    src = img_tag.get('src')
    if not src:
        return
    img_path = (EPUB_DIR / src).resolve()
    if not img_path.exists():
        print('!! image not found:', src)
        return
    try:
        doc.add_picture(str(img_path), width=Inches(5.5))
    except Exception as e:
        print('!! failed to embed image', src, e)
    alt = img_tag.get('alt')
    if alt:
        cap = doc.add_paragraph()
        cap.alignment = 1
        add_plain_run(cap, alt, bold=False)


def process_block(doc, tag):
    if is_hidden(tag):
        return
    name = tag.name
    if name in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
        level = min(int(name[1]), 4)
        para = doc.add_heading(level=level)
        walk_inline(para, tag)
        for run in para.runs:
            set_run_font(run, size_half_points=max(BASE_SZ, 32 - level * 4), bold=True)
    elif name == 'p':
        cls = tag.get('class') or []
        para = doc.add_paragraph()
        if any(c == 'note' or c.endswith('-note') for c in cls):
            add_plain_run(para, '※')
        walk_inline(para, tag)
    elif name == 'blockquote':
        handled_children = False
        for child in tag.find_all(recursive=False):
            if is_hidden(child):
                continue
            if child.name == 'p':
                para = doc.add_paragraph()
                para.paragraph_format.left_indent = Pt(24)
                walk_inline(para, child)
                handled_children = True
        if not handled_children:
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Pt(24)
            walk_inline(para, tag)
    elif name in ('ul', 'ol'):
        for li in tag.find_all('li', recursive=False):
            if is_hidden(li):
                continue
            para = doc.add_paragraph(style='List Bullet')
            walk_inline(para, li)
    elif name == 'figure':
        img = tag.find('img')
        if img is not None:
            add_image(doc, img)
        cap = tag.find('figcaption')
        if cap is not None:
            para = doc.add_paragraph()
            para.alignment = 1
            walk_inline(para, cap)
    else:
        for child in tag.find_all(recursive=False):
            process_block(doc, child)


def main():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = BASE_FONT
    style.font.size = Pt(BASE_SZ / 2)
    files, book_title = spine_files()
    title = doc.add_heading(f'{book_title}（全文・ルビ付き）', level=0)
    for run in title.runs:
        set_run_font(run, size_half_points=40, bold=True)
    for fname in files:
        path = EPUB_DIR / fname
        if not path.exists():
            print('missing', fname)
            continue
        soup = BeautifulSoup(path.read_text(encoding='utf-8'), 'lxml-xml')
        body = soup.find('body')
        if body is None:
            continue
        for child in body.find_all(recursive=False):
            process_block(doc, child)
        doc.add_page_break()
    doc.save(OUT_PATH)
    print('saved', OUT_PATH)


if __name__ == '__main__':
    main()
